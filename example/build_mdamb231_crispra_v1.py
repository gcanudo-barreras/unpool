"""ONE SOURCE, ONE BUILDER. Written from `templates/builder.py`.

================================================================================
MDA-MB-231 CRISPR-activation orthotopic xenograft -- 5 arms, 41 mice, from
Galal et al. 2026, Nature Communications, DOI 10.1038/s41467-026-76293-x.
================================================================================

THE SOURCE
------------------------------------------------------------------------------
File: `41467_2026_76293_MOESM6_ESM.xlsx` (the article's Source Data file),
sheet `Supplementary Fig. 8b.` -- the trailing period is part of the sheet
name -- header rows 3-4, data rows 5-13, columns C..AQ.
Obtained from: Europe PMC supplementary bundle for PMC13534483.
  sha256 of the workbook:
  13fbff4ba57d7de064911c288f0002757754dbfe26f45eea8bffc94b8d4ebae7
Licence: CC BY-NC-ND 4.0 -- travels with the data, into the build report.
  NC and ND are real terms, and they are the reason this repository ships the
  BUILDER rather than the rows. Re-run this file against the publisher's own
  bytes and you have the CSV; it is not redistributed from here.

MODEL, HOST AND METHOD, verified against the paper's own text
------------------------------------------------------------------------------
  * METHOD, stated explicitly: "Primary tumor volumes were assessed regularly
    using the formula: [4/3 x pi x (length/2) x (width/2)2]" -- caliper
    measurement, Methods, "In vivo orthotopic xenograft studies (mammary fat
    pad transplantation)". Volumes are mm3; the sheet's own banner says so.
  * MODEL: MDA-MB-231, taken from the paper's text and not from the sheet's
    title: "MDA-MB-231 cells transduced with the endogenously activated
    targets of CRISPR lentiSAMv2 cloned constructs (1 sgRNA/gene) ... 2 x
    10(6) cells/mouse were inoculated into the MFP of NSG female mice".
    Orthotopic mammary fat pad, female NSG host. The paper identifies the line
    by Cellosaurus accession CVCL_0062.
  * WHICH ARM IS THE CONTROL: `NTC`, and the sheet never expands it. The paper
    does: "non-targeting control (NTC)". One arm, not two pooled.
  * GROUP SIZES: the paper states "n = 8-9 mice/group". The data carries
    NTC 9, VPS45g1a 8, RBSNg2a 8, CMTR2g1a 8, NF2g1a 8 = 41. These AGREE, and
    the data settles what the paper leaves open -- which group holds the 9.
    Both numbers are written down, because agreement is only visible if you
    print both.
  * ENDPOINT: "The injected mice were euthanized on day 56 of inoculation",
    and the last column of the grid is day 56. Consistent.
  * ATTRITION: the paper states no exclusions or deaths for this experiment.
    The sheet is a complete 41 x 9 rectangle -- no blank cell, no censored
    animal -- which is consistent with none. The protocol's welfare cap is
    "a maximum permitted tumor volume of 2000 mm3"; the largest value anywhere
    on the sheet is 932.34 mm3, so no animal approached it.

WHAT A HUMAN RESOLVED THAT THE PAPER DOES NOT STATE
------------------------------------------------------------------------------
**This section is the reason the file exists.**

  1. `Mouse #1` IS NOT AN IDENTIFIER. The column labels restart at `Mouse #1`
     inside every one of the five arms, so five different animals share each
     label. This build namespaces them `MDA-MB-231:<arm>:<index>`.
     **If you key on the sheet's own label you fuse five mice into one and the
     build emits 9 animals instead of 41** -- silently, with a plausible CSV.
     The index is POSITIONAL and ours; the authors assigned no subject IDs.

  2. THE DAY-0 ROW IS NOT A MEASUREMENT. Every one of the 41 mice is recorded
     as 0 mm3 on day 0. A caliper cannot measure a 0 mm3 tumour; that row is
     the inoculation baseline written as a zero. It is KEPT, because it is what
     the authors published and dropping it is an unstated derivative decision.
     **Anything downstream that takes a log or a ratio must handle it.** The
     scanner already does, implicitly: its own grid for this sheet is 31..56
     and its `fold=2.7` is day 31 -> day 56, not day 0 -> day 56.

  3. THE GUIDE SUFFIXES ARE NOT RESOLVED. Four arms are labelled `g1a` and
     `RBSN` is `g2a`. The paper says "1 sgRNA/gene" and never maps a guide
     number to an arm, so which sgRNA each arm carries is NOT derivable from
     the documents and is not claimed here. The arm label is carried through
     verbatim instead of being interpreted. (`RBSN` is `ZFYVE20`, which is the
     name the paper's knockout section uses for the same gene.)

  4. THE `study` LABEL IS OURS. `MDA-MB-231-CRISPRa-MFP` is a key this build
     invents to name the experiment. The authors wrote no such string.

THE SUBJECT KEY
------------------------------------------------------------------------------
`MDA-MB-231:<arm>:<n>`, POSITIONAL AND SYNTHETIC. `<n>` is the mouse number
printed in the sheet's column header, which is unique only within its arm. No
identifier in this CSV was assigned by the authors.
"""
import csv
import io
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from unpool.guard import assert_buildable, CorpusGuardRefused  # noqa: E402

_HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(os.path.dirname(_HERE), "si_cache", "PMC13534483.zip")
MEMBER = "41467_2026_76293_MOESM6_ESM.xlsx"
SHEET = "Supplementary Fig. 8b."
OUT = os.path.join(_HERE, "corpus", "mdamb231_crispra_v1.csv")

SOURCE = "Galal-NCOMMS76293"
MODEL = "MDA-MB-231"
STUDY = "MDA-MB-231-CRISPRa-MFP"
CONTROL_ARM = "NTC"

# THE ROUND-TRIP GUARD. Counted BY READING THE SHEET, not copied from the
# legend: the paper says "n = 8-9 mice/group", which is a range and not a
# check. If a re-download disagrees with these, the build REFUSES. It does not
# adapt -- a source that changed under you is a fact to investigate.
_VERIFIED_N = 41
_VERIFIED_GRID = (0.0, 31.0, 34.0, 38.0, 41.0, 44.0, 48.0, 51.0, 56.0)
_VERIFIED_ARMS = {"NTC": 9, "VPS45g1a": 8, "RBSNg2a": 8,
                  "CMTR2g1a": 8, "NF2g1a": 8}
_VERIFIED_SHA256 = ("13fbff4ba57d7de064911c288f0002757754dbfe26f45eea8"
                    "bffc94b8d4ebae7")


class BuildRefused(Exception):
    """The source does not match what was verified. Stop."""


def _workbook_bytes():
    if not os.path.exists(SRC):
        raise BuildRefused(
            "%s not found. The supplementary bundle is not in the repository "
            "(si_cache/ is gitignored); fetch it, which takes about 18 MB:\n"
            "    python -c \"from unpool import scan; print(scan.fetch('PMC13534483'))\"\n"
            "A missing input is a wall, not an empty result (protocol rule "
            "10)." % SRC)
    with zipfile.ZipFile(SRC) as z:
        if MEMBER not in z.namelist():
            raise BuildRefused("%s holds no member %s" % (SRC, MEMBER))
        return z.read(MEMBER)


def read_source():
    """-> {animal_id: (arm, [(day, volume), ...])}  Parse and nothing else."""
    import hashlib
    try:
        import openpyxl
    except ImportError as exc:      # protocol rule 10: say what blocked you
        raise BuildRefused(
            "openpyxl is not importable in this interpreter (%s). That is a "
            "missing parser, NOT an empty source. `pip install openpyxl`, and "
            "check that you are running the interpreter you think you are."
            % exc)

    blob = _workbook_bytes()
    got = hashlib.sha256(blob).hexdigest()
    if got != _VERIFIED_SHA256:
        raise BuildRefused(
            "%s sha256 %s, verified %s -- the publisher's file changed under "
            "this build. Investigate before editing this constant."
            % (MEMBER, got, _VERIFIED_SHA256))

    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True,
                                data_only=True)
    if SHEET not in wb.sheetnames:
        raise BuildRefused("expected sheet %r, have %s"
                           % (SHEET, wb.sheetnames))
    rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]

    arm_row, head_row, data = rows[2], rows[3], rows[4:]
    days = [float(r[1]) for r in data]

    curves, arm = {}, None
    for col, label in enumerate(head_row):
        if not (isinstance(label, str) and label.strip().startswith("Mouse")):
            continue
        # The arm name is written once, over the first column of its block.
        banner = arm_row[col]
        if isinstance(banner, str) and banner.strip():
            arm = banner.strip()
        if arm is None:
            raise BuildRefused("column %d carries a mouse label under no arm "
                               "banner" % col)
        # `Mouse #1` restarts inside every arm -- see the docstring. The arm is
        # part of the key, or five animals silently become one.
        n = label.strip().replace("Mouse #", "").strip()
        aid = "%s:%s:%s" % (MODEL, arm, n)
        if aid in curves:
            raise BuildRefused("duplicate subject key %s" % aid)
        vals = [r[col] for r in data]
        if any(not isinstance(v, (int, float)) for v in vals):
            raise BuildRefused(
                "%s has a non-numeric or empty cell; this sheet was verified "
                "as a complete rectangle, so a gap means the source changed "
                "or the parse did" % aid)
        curves[aid] = (arm, list(zip(days, [float(v) for v in vals])))
    return curves


def main():
    curves = read_source()

    if len(curves) != _VERIFIED_N:
        raise BuildRefused(
            "%d subjects, verified %d -- the source changed, or the parse did. "
            "Investigate before editing this constant."
            % (len(curves), _VERIFIED_N))

    per_arm = {}
    for arm, _ in curves.values():
        per_arm[arm] = per_arm.get(arm, 0) + 1
    if per_arm != _VERIFIED_ARMS:
        raise BuildRefused("arm sizes %s, verified %s"
                           % (sorted(per_arm.items()),
                              sorted(_VERIFIED_ARMS.items())))

    grids = {tuple(c for c, _ in pts) for _, pts in curves.values()}
    if grids != {_VERIFIED_GRID}:
        raise BuildRefused("covariate grid %s, verified %s"
                           % (sorted(grids), [_VERIFIED_GRID]))

    rows = [{"animal_id": aid, "model": MODEL, "study": STUDY,
             "arm_label": arm,
             "role": "control" if arm == CONTROL_ARM else "treated",
             "source": SOURCE, "days": "%g" % d, "volumes": "%g" % v}
            for aid, (arm, pts) in sorted(curves.items()) for d, v in pts]

    # Refuse before writing, never after: a guard that runs on a file already
    # on disk is a report, not a gate.
    n_keys = assert_buildable(rows, label=OUT)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), delimiter=";")
        w.writeheader()
        w.writerows(rows)
    print("wrote %s: %d rows, %d subjects, %d (model, study) key(s)"
          % (OUT, len(rows), len(curves), n_keys))
    for arm in sorted(per_arm):
        print("    %-10s n=%-3d %s"
              % (arm, per_arm[arm], "CONTROL" if arm == CONTROL_ARM else ""))


if __name__ == "__main__":
    try:
        main()
    except (BuildRefused, CorpusGuardRefused) as exc:
        print("REFUSED\n%s" % exc)
        sys.exit(1)
